"""
Convert Markdown files to Word documents
Requires: pip install python-docx markdown
"""

import os
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import markdown
    from html.parser import HTMLParser
except ImportError:
    print("Installing required packages...")
    import subprocess
    subprocess.check_call(['pip', 'install', 'python-docx', 'markdown'])
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import markdown
    from html.parser import HTMLParser


class MarkdownToWordConverter:
    """Convert Markdown to Word document with formatting"""
    
    def __init__(self):
        self.doc = Document()
        
    def convert_file(self, md_file, output_file=None):
        """Convert a markdown file to Word document"""
        if output_file is None:
            output_file = md_file.replace('.md', '.docx')
        
        print(f"Converting {md_file} to {output_file}...")
        
        # Read markdown file
        with open(md_file, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # Convert markdown to HTML
        html = markdown.markdown(md_content, extensions=['tables', 'fenced_code'])
        
        # Parse and add to document
        self._parse_html(html)
        
        # Save document
        self.doc.save(output_file)
        print(f"✅ Successfully created: {output_file}")
        
    def _parse_html(self, html):
        """Parse HTML and add formatted content to Word document"""
        lines = html.split('\n')
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
                
            # Headers
            if line.startswith('<h1>'):
                text = self._clean_html(line)
                heading = self.doc.add_heading(text, level=1)
                heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif line.startswith('<h2>'):
                text = self._clean_html(line)
                self.doc.add_heading(text, level=2)
            elif line.startswith('<h3>'):
                text = self._clean_html(line)
                self.doc.add_heading(text, level=3)
            elif line.startswith('<h4>'):
                text = self._clean_html(line)
                self.doc.add_heading(text, level=4)
            
            # Paragraphs
            elif line.startswith('<p>'):
                text = self._clean_html(line)
                if text:
                    self.doc.add_paragraph(text)
            
            # Lists
            elif line.startswith('<li>'):
                text = self._clean_html(line)
                self.doc.add_paragraph(text, style='List Bullet')
            
            # Code blocks
            elif line.startswith('<code>') or line.startswith('<pre>'):
                text = self._clean_html(line)
                p = self.doc.add_paragraph(text)
                p.style = 'No Spacing'
                run = p.runs[0] if p.runs else p.add_run()
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
            
            # Horizontal rule
            elif line.startswith('<hr'):
                self.doc.add_paragraph('_' * 50)
    
    def _clean_html(self, html_text):
        """Remove HTML tags and decode entities"""
        # Remove tags
        text = html_text
        tags = ['<h1>', '</h1>', '<h2>', '</h2>', '<h3>', '</h3>', '<h4>', '</h4>',
                '<p>', '</p>', '<li>', '</li>', '<strong>', '</strong>', '<em>', '</em>',
                '<code>', '</code>', '<pre>', '</pre>', '<ul>', '</ul>', '<ol>', '</ol>',
                '<table>', '</table>', '<tr>', '</tr>', '<td>', '</td>', '<th>', '</th>']
        
        for tag in tags:
            text = text.replace(tag, '')
        
        # Decode common HTML entities
        text = text.replace('&lt;', '<')
        text = text.replace('&gt;', '>')
        text = text.replace('&amp;', '&')
        text = text.replace('&quot;', '"')
        text = text.replace('&#39;', "'")
        
        return text.strip()


def main():
    """Convert all documentation markdown files to Word"""
    
    # Define files to convert
    files_to_convert = [
        ('FRAMEWORK_DOCUMENTATION.md', 'WinETL_InfinityX_Technical_Documentation.docx'),
        ('QUICKSTART_GUIDE.md', 'WinETL_InfinityX_QuickStart_Guide.docx'),
        ('PUBLISHING_GUIDE.md', 'WinETL_InfinityX_Publishing_Guide.docx'),
    ]
    
    base_path = Path(__file__).parent
    
    print("=" * 60)
    print("📄 Converting Markdown Documentation to Word Documents")
    print("=" * 60)
    print()
    
    for md_file, docx_file in files_to_convert:
        md_path = base_path / md_file
        docx_path = base_path / docx_file
        
        if not md_path.exists():
            print(f"⚠️  Warning: {md_file} not found, skipping...")
            continue
        
        try:
            converter = MarkdownToWordConverter()
            converter.convert_file(str(md_path), str(docx_path))
        except Exception as e:
            print(f"❌ Error converting {md_file}: {e}")
    
    print()
    print("=" * 60)
    print("✅ Conversion Complete!")
    print("=" * 60)
    print()
    print("📁 Word documents created in:")
    print(f"   {base_path}")
    print()


if __name__ == "__main__":
    main()
